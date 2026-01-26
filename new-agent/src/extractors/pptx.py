"""
PowerPoint content extractor using python-pptx.
"""

from __future__ import annotations

from pathlib import Path
from typing import List

from src.extractors.base import ContentExtractor, ExtractionResult


class PPTXExtractor(ContentExtractor):
    """Extract text content from PowerPoint files."""
    
    @property
    def supported_extensions(self) -> List[str]:
        return [".pptx", ".ppt"]
    
    def extract(self, file_path: Path) -> ExtractionResult:
        """Extract text from PowerPoint file."""
        
        # Handle .ppt (old format) - just return a note
        if file_path.suffix.lower() == ".ppt":
            return ExtractionResult(
                file_path=str(file_path),
                file_type=".ppt",
                success=False,
                error="Legacy .ppt format not supported. Please convert to .pptx",
            )
        
        try:
            from pptx import Presentation
        except ImportError:
            return ExtractionResult(
                file_path=str(file_path),
                file_type=".pptx",
                success=False,
                error="python-pptx not installed. Run: pip install python-pptx",
            )
        
        try:
            prs = Presentation(file_path)
            
            text_parts = []
            for slide_num, slide in enumerate(prs.slides, start=1):
                slide_texts = []
                
                # Get slide title if present
                if slide.shapes.title and slide.shapes.title.text:
                    slide_texts.append(f"Title: {slide.shapes.title.text}")
                
                # Get text from all shapes
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        # Skip the title we already got
                        if shape == slide.shapes.title:
                            continue
                        slide_texts.append(shape.text)
                    
                    # Also get text from tables
                    if shape.has_table:
                        table = shape.table
                        for row in table.rows:
                            row_text = " | ".join(
                                cell.text for cell in row.cells if cell.text
                            )
                            if row_text:
                                slide_texts.append(row_text)
                
                # Get notes if present
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    notes = slide.notes_slide.notes_text_frame.text
                    if notes and notes.strip():
                        slide_texts.append(f"[Notes: {notes.strip()}]")
                
                if slide_texts:
                    text_parts.append(f"--- Slide {slide_num} ---\n" + "\n".join(slide_texts))
            
            return ExtractionResult(
                file_path=str(file_path),
                file_type=".pptx",
                success=True,
                text_content="\n\n".join(text_parts),
                page_count=len(prs.slides),
            )
            
        except Exception as e:
            return ExtractionResult(
                file_path=str(file_path),
                file_type=".pptx",
                success=False,
                error=f"Failed to extract PowerPoint: {e}",
            )
