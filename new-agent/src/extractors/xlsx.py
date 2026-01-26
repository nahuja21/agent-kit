"""
Excel content extractor using openpyxl.
"""

from __future__ import annotations

from pathlib import Path
from typing import List

from src.extractors.base import ContentExtractor, ExtractionResult


class XLSXExtractor(ContentExtractor):
    """Extract text content from Excel files."""
    
    @property
    def supported_extensions(self) -> List[str]:
        return [".xlsx", ".xls"]
    
    def extract(self, file_path: Path) -> ExtractionResult:
        """Extract text from Excel file."""
        
        # Handle .xls (old format)
        if file_path.suffix.lower() == ".xls":
            return ExtractionResult(
                file_path=str(file_path),
                file_type=".xls",
                success=False,
                error="Legacy .xls format not supported. Please convert to .xlsx",
            )
        
        try:
            from openpyxl import load_workbook
        except ImportError:
            return ExtractionResult(
                file_path=str(file_path),
                file_type=".xlsx",
                success=False,
                error="openpyxl not installed. Run: pip install openpyxl",
            )
        
        try:
            # Load workbook in read-only mode for better performance
            wb = load_workbook(file_path, read_only=True, data_only=True)
            
            text_parts = []
            sheet_count = 0
            
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                sheet_count += 1
                
                sheet_texts = [f"\n--- Sheet: {sheet_name} ---"]
                
                # Get dimensions
                max_row = min(sheet.max_row or 0, 100)  # Limit to first 100 rows
                max_col = min(sheet.max_column or 0, 20)  # Limit to first 20 columns
                
                if max_row == 0 or max_col == 0:
                    continue
                
                # Extract data as table format
                rows_extracted = 0
                for row in sheet.iter_rows(min_row=1, max_row=max_row, max_col=max_col, values_only=True):
                    # Convert values to strings, handle None
                    row_values = [str(cell) if cell is not None else "" for cell in row]
                    row_text = " | ".join(row_values)
                    
                    # Skip empty rows
                    if any(v.strip() for v in row_values):
                        sheet_texts.append(row_text)
                        rows_extracted += 1
                
                if rows_extracted > 0:
                    if max_row >= 100:
                        sheet_texts.append(f"[... showing first 100 rows ...]")
                    text_parts.extend(sheet_texts)
            
            wb.close()
            
            return ExtractionResult(
                file_path=str(file_path),
                file_type=".xlsx",
                success=True,
                text_content="\n".join(text_parts),
                metadata={"sheet_count": str(sheet_count)},
            )
            
        except Exception as e:
            return ExtractionResult(
                file_path=str(file_path),
                file_type=".xlsx",
                success=False,
                error=f"Failed to extract Excel: {e}",
            )
