"""
Word document content extractor using python-docx.
"""

from __future__ import annotations

from pathlib import Path
from typing import List

from src.extractors.base import ContentExtractor, ExtractionResult


class DOCXExtractor(ContentExtractor):
    """Extract text content from Word documents."""
    
    @property
    def supported_extensions(self) -> List[str]:
        return [".docx", ".doc"]
    
    def extract(self, file_path: Path) -> ExtractionResult:
        """Extract text from Word document."""
        
        # Handle .doc (old format)
        if file_path.suffix.lower() == ".doc":
            return ExtractionResult(
                file_path=str(file_path),
                file_type=".doc",
                success=False,
                error="Legacy .doc format not supported. Please convert to .docx",
            )
        
        try:
            from docx import Document
        except ImportError:
            return ExtractionResult(
                file_path=str(file_path),
                file_type=".docx",
                success=False,
                error="python-docx not installed. Run: pip install python-docx",
            )
        
        try:
            doc = Document(file_path)
            
            text_parts = []
            
            # Get document properties if available
            metadata = {}
            if doc.core_properties:
                if doc.core_properties.title:
                    metadata["title"] = doc.core_properties.title
                if doc.core_properties.author:
                    metadata["author"] = doc.core_properties.author
            
            # Extract paragraphs
            for para in doc.paragraphs:
                if para.text.strip():
                    # Check if it's a heading
                    if para.style and para.style.name.startswith("Heading"):
                        text_parts.append(f"\n## {para.text}\n")
                    else:
                        text_parts.append(para.text)
            
            # Extract tables
            for table_num, table in enumerate(doc.tables, start=1):
                table_texts = [f"\n[Table {table_num}]"]
                for row in table.rows:
                    row_text = " | ".join(
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    )
                    if row_text:
                        table_texts.append(row_text)
                if len(table_texts) > 1:  # More than just the header
                    text_parts.extend(table_texts)
            
            return ExtractionResult(
                file_path=str(file_path),
                file_type=".docx",
                success=True,
                text_content="\n".join(text_parts),
                metadata=metadata,
            )
            
        except Exception as e:
            return ExtractionResult(
                file_path=str(file_path),
                file_type=".docx",
                success=False,
                error=f"Failed to extract Word document: {e}",
            )
